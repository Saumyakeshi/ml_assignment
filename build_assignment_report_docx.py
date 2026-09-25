from __future__ import annotations

import re
import sys
from pathlib import Path
from urllib.parse import quote


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "tmp" / "docdeps"))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SOURCE = ROOT / "reports" / "drafts" / "assignment-report-4000-words.md"
OUTPUT = ROOT / "output" / "documents" / "machine-learning-in-cyber-assignment-report.docx"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
GREY = "666666"
LIGHT_GREY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"


def clean_text(text: str) -> str:
    """Use plain, portable punctuation in the generated report."""
    return (
        text.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
        .replace("\u00a0", " ")
        .replace("\ufffd", "-")
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LIGHT_GREY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_keep_together(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepLines")
    p_pr.append(keep)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_begin, instr, fld_sep, fld_text, fld_end))
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREY)


def add_hyperlink(paragraph, text: str, target: str, *, bold=False, italic=False):
    target_path = (SOURCE.parent / clean_text(target)).resolve()
    address = target_path.as_uri() if target_path.exists() else target
    part = paragraph.part
    rel_id = part.relate_to(
        address,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    if bold:
        run_props.append(OxmlElement("w:b"))
    if italic:
        run_props.append(OxmlElement("w:i"))
    new_run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = clean_text(text)
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\[[^\]]+\]\[[^\]]+\]|\[(?:L\d{2}|[A-Z][A-Z0-9]*)\])"
)


def add_inline(paragraph, text: str, ref_defs: dict[str, str]) -> None:
    text = clean_text(text)
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        elif "](" in token:
            label, target = token[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, target)
        elif "][" in token:
            label, ref = token[1:-1].split("][", 1)
            add_hyperlink(paragraph, label, ref_defs.get(ref, ref))
        else:
            run = paragraph.add_run(token)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_body_paragraph(doc: Document, text: str, ref_defs: dict[str, str], style=None):
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text, ref_defs)
    set_keep_together(paragraph)
    return paragraph


def add_metric_table(doc: Document, rows: list[list[str]], ref_defs: dict[str, str]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, value in enumerate(row_values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            add_inline(paragraph, value, ref_defs)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph()


FIGURES = [
    (
        "Email security",
        ROOT / "section_01_results/reports/figures/data-exploration.png",
        "Email data exploration: class balance and normalized message length.",
    ),
    (
        "Email security",
        ROOT / "section_01_results/reports/figures/classic-evaluation.png",
        "Logistic Regression test evaluation.",
    ),
    (
        "Email security",
        ROOT / "section_01_results/reports/figures/lstm-evaluation.png",
        "Email LSTM test evaluation.",
    ),
    (
        "Email security",
        ROOT / "section_01_results/reports/figures/lstm-training-history.png",
        "Email LSTM training and validation history.",
    ),
    (
        "Network intrusion",
        ROOT / "reports/section_02/figures/random-forest-evaluation.png",
        "Random Forest intrusion evaluation (retained legacy evidence).",
    ),
    (
        "Network intrusion",
        ROOT / "reports/section_02/figures/cnn-evaluation.png",
        "CNN intrusion evaluation (retained legacy evidence).",
    ),
    (
        "Traffic anomalies",
        ROOT / "reports/section_03/figures/isolation-forest-evaluation.png",
        "Isolation Forest anomaly evaluation (retained legacy evidence).",
    ),
    (
        "Traffic anomalies",
        ROOT / "reports/section_03/figures/autoencoder-evaluation.png",
        "Autoencoder anomaly evaluation (retained legacy evidence).",
    ),
    (
        "Ransomware",
        ROOT / "reports/section_04/figures/svm-evaluation.png",
        "SVM ransomware evaluation (retained legacy evidence).",
    ),
    (
        "Ransomware",
        ROOT / "reports/section_04/figures/lstm-evaluation.png",
        "LSTM ransomware evaluation (retained legacy evidence).",
    ),
]


def set_picture_alt_text(paragraph, description: str) -> None:
    drawings = paragraph._p.xpath(".//wp:docPr")
    if drawings:
        drawings[-1].set("descr", description)
        drawings[-1].set("title", description)


def add_figures(doc: Document) -> None:
    for index, (section_name, image_path, caption) in enumerate(FIGURES, start=1):
        if index == 1 or index % 2 == 1:
            doc.add_page_break()
            heading = doc.add_paragraph(section_name, style="Figure Group")
            set_keep_with_next(heading)
        if not image_path.exists():
            warning = doc.add_paragraph(f"Figure unavailable: {image_path}")
            warning.runs[0].font.color.rgb = RGBColor.from_string("C00000")
            continue
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(6.25))
        set_picture_alt_text(paragraph, caption)
        caption_paragraph = doc.add_paragraph(style="Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.add_run(f"Figure {index}. {caption}")
        set_keep_together(caption_paragraph)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(12)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(GREY)
    subtitle.paragraph_format.space_after = Pt(18)

    heading1 = styles["Heading 1"]
    heading1.font.name = "Arial"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor.from_string(BLACK)
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(7)
    heading1.paragraph_format.keep_with_next = True

    heading2 = styles["Heading 2"]
    heading2.font.name = "Arial"
    heading2.font.size = Pt(12)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor.from_string(BLACK)
    heading2.paragraph_format.space_before = Pt(10)
    heading2.paragraph_format.space_after = Pt(5)
    heading2.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GREY)
    caption.paragraph_format.space_after = Pt(9)

    figure_group = styles.add_style("Figure Group", WD_STYLE_TYPE.PARAGRAPH)
    figure_group.font.name = "Arial"
    figure_group.font.size = Pt(11)
    figure_group.font.bold = True
    figure_group.font.color.rgb = RGBColor.from_string(BLACK)
    figure_group.paragraph_format.space_after = Pt(5)

    contents = styles.add_style("Contents Entry", WD_STYLE_TYPE.PARAGRAPH)
    contents.font.name = "Arial"
    contents.font.size = Pt(10.5)
    contents.font.color.rgb = RGBColor.from_string(NAVY)
    contents.paragraph_format.left_indent = Cm(0.4)
    contents.paragraph_format.first_line_indent = Cm(-0.4)
    contents.paragraph_format.space_after = Pt(4)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)


def add_header_footer(section) -> None:
    header = section.header.paragraphs[0]
    header.text = "Machine Learning in Cyber - Assignment Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_number(section.footer.paragraphs[0])


def parse_source() -> tuple[list[str], dict[str, str]]:
    raw_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    ref_defs: dict[str, str] = {}
    content: list[str] = []
    for line in raw_lines:
        match = re.fullmatch(r"\[([^\]]+)\]:\s+(.+)", line.strip())
        if match:
            ref_defs[match.group(1)] = match.group(2)
        else:
            content.append(line)
    return content, ref_defs


def add_title_page(doc: Document, ref_defs: dict[str, str]) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(45)
    title = doc.add_paragraph("Machine Learning in Cyber", style="Title")
    title.add_run("\nImplementation and Critical Evaluation")
    subtitle = doc.add_paragraph("Expanded assignment report", style="Subtitle")
    subtitle.add_run("\nUpdated 25 September 2026")

    meta = doc.add_table(rows=2, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta.columns[0].width = Cm(4.2)
    meta.columns[1].width = Cm(11.5)
    labels = (("Main-text word count", "4,000"), ("Evidence base", "Four cybersecurity machine-learning experiments"))
    for row, (label, value) in zip(meta.rows, labels):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            set_cell_margins(cell, top=80, start=60, bottom=80, end=60)
    set_table_borders(meta, color=WHITE, size="0")

    doc.add_paragraph()
    warning = doc.add_paragraph()
    warning.add_run("Submission warning. ").bold = True
    warning.add_run(
        "The extracted assignment brief specifies a maximum of 3,000 words. "
        "This expanded 4,000-word draft must be shortened, or an authoritative amendment obtained, before submission."
    )
    warning.paragraph_format.space_before = Pt(8)

    evidence = doc.add_paragraph()
    evidence.add_run("Evidence policy. ").bold = True
    evidence.add_run(
        "Section 1 uses the fresh export produced by the updated notebook on 25 September 2026. "
        "Sections 2-4 use retained JSON evidence from legacy executions and require fresh reruns before their revised notebooks can replace those results. "
    )
    add_hyperlink(evidence, "The alignment review", ref_defs["A"])
    evidence.add_run(" records the remaining provenance limitations.")

    doc.add_page_break()
    doc.add_paragraph("Contents", style="Heading 1")
    entries = [
        "1. Purpose and relationship to the lectures",
        "2. Experimental design and evaluation principles",
        "3. Email security: Logistic Regression and LSTM",
        "4. Network intrusion: Random Forest and CNN",
        "5. Traffic anomalies: Isolation Forest and Autoencoder",
        "6. Ransomware: API n-grams and sequence learning",
        "7. Cross-experiment assessment and conclusion",
        "Evidence appendix",
        "References and implementation evidence",
    ]
    for entry in entries:
        doc.add_paragraph(entry, style="Contents Entry")
    doc.add_page_break()


def build_document() -> None:
    lines, ref_defs = parse_source()
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_header_footer(doc.sections[0])
    add_title_page(doc, ref_defs)

    index = 0
    paragraph_buffer: list[str] = []
    skipped_front_matter = True

    def flush_buffer() -> None:
        if paragraph_buffer:
            add_body_paragraph(doc, " ".join(part.strip() for part in paragraph_buffer), ref_defs)
            paragraph_buffer.clear()

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if line == "<!-- report-body:start -->":
            skipped_front_matter = False
            index += 1
            continue
        if skipped_front_matter:
            index += 1
            continue
        if line in ("<!-- report-body:end -->", ""):
            flush_buffer()
            index += 1
            continue
        if line.startswith("<!--"):
            index += 1
            continue

        if line.startswith("## "):
            flush_buffer()
            title = clean_text(line[3:])
            if title.startswith("Evidence appendix") or title.startswith("References and implementation"):
                doc.add_page_break()
            heading = doc.add_paragraph(title, style="Heading 1")
            set_keep_with_next(heading)
            index += 1
            continue

        if line.startswith("### "):
            flush_buffer()
            title = clean_text(line[4:])
            heading = doc.add_paragraph(title, style="Heading 2")
            set_keep_with_next(heading)
            if title == "Existing figures":
                index += 1
                explanatory: list[str] = []
                while index < len(lines):
                    current = lines[index].strip()
                    if current.startswith("## "):
                        break
                    if current and not current.startswith("-"):
                        explanatory.append(current)
                    index += 1
                if explanatory:
                    add_body_paragraph(doc, " ".join(explanatory), ref_defs)
                add_figures(doc)
            else:
                index += 1
            continue

        if line.startswith("|"):
            flush_buffer()
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                add_metric_table(doc, rows, ref_defs)
            continue

        bullet = re.match(r"^-\s+(.+)", line)
        numbered = re.match(r"^\d+\.\s+(.+)", line)
        if bullet or numbered:
            flush_buffer()
            content = (bullet or numbered).group(1)
            style = "List Bullet" if bullet else "List Number"
            add_body_paragraph(doc, content, ref_defs, style=style)
            index += 1
            continue

        paragraph_buffer.append(line)
        index += 1

    flush_buffer()

    # Repeat page setup in case Word introduces additional sections later.
    configure_sections(doc)
    for section in doc.sections:
        add_header_footer(section)

    core = doc.core_properties
    core.title = "Machine Learning in Cyber: Implementation and Critical Evaluation"
    core.subject = "Cybersecurity machine-learning assignment report"
    core.keywords = "machine learning, cybersecurity, classification, anomaly detection, evaluation"
    core.comments = "Generated from the updated assignment report and repository evidence."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
