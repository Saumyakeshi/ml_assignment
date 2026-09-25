from __future__ import annotations

import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "tmp" / "docdeps"))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = ROOT / "output" / "documents" / "section-1-email-security-report.docx"

BLACK = "000000"
BLUE = "1F4E78"
PALE_BLUE = "EEF5FA"
LIGHT_GREY = "D9D9D9"
GREY = "666666"
WHITE = "FFFFFF"


def set_font(run, name="Arial", size=11, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GREY)
        borders.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, display, end))
    set_font(run, size=9, color=GREY)


def add_hyperlink(paragraph, label, relative_path):
    target = (ROOT / relative_path).resolve()
    address = target.as_uri()
    rel_id = paragraph.part.relate_to(
        address,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend((color, underline))
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc, text, *, bold_lead=None, style=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_subsection_heading(doc, text):
    page_start = doc.add_paragraph()
    page_start.paragraph_format.page_break_before = True
    page_start.paragraph_format.space_after = Pt(12)
    spacer = page_start.add_run(" ")
    set_font(spacer, size=4)
    heading = doc.add_paragraph(text, style="Heading 2")
    return heading


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_metrics_table(doc, rows):
    headers = ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC AUC", "Average precision"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            if row_index % 2 == 1:
                shade_cell(cell, PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            set_font(run, size=8.5)
    doc.add_paragraph()


def add_confusion_table(doc):
    headers = ["Model", "True negatives", "False positives", "False negatives", "True positives"]
    rows = [
        ["Logistic Regression", "369", "2", "5", "67"],
        ["LSTM", "371", "0", "26", "46"],
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    repeat_header(table.rows[0])
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, BLUE)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=8.5, bold=True, color=WHITE)
    for r_index, values in enumerate(rows, start=1):
        row = table.add_row()
        for c_index, value in enumerate(values):
            cell = row.cells[c_index]
            if r_index % 2 == 1:
                shade_cell(cell, PALE_BLUE)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(value), size=8.5)
    doc.add_paragraph()


def add_figure(doc, image_path, caption, number, width=6.35, page_break_before=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.page_break_before = page_break_before
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    doc_pr = paragraph._p.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[-1].set("descr", caption)
        doc_pr[-1].set("title", caption)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Figure {number}. {caption}")
    cap.paragraph_format.keep_with_next = False


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(GREY)
    subtitle.paragraph_format.space_after = Pt(20)

    h1 = styles["Heading 1"]
    h1.font.name = "Arial"
    h1.font.size = Pt(17)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLACK)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Arial"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLACK)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GREY)
    caption.paragraph_format.space_after = Pt(8)

    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    page_number(section.first_page_footer.paragraphs[0])

    header = section.header.paragraphs[0]
    header.text = "Section 1 Email Security and Phishing Detection"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=8, color=GREY)
    page_number(section.footer.paragraphs[0])


def cover_page(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(50)
    doc.add_paragraph("Machine Learning in Cyber", style="Title")
    doc.add_paragraph("Section 1 Email Security and Phishing Detection", style="Subtitle")

    details = [
        ("Module", "COMP70049"),
        ("Assessment", "Machine Learning in Cyber Assignment"),
        ("Student name", "____________________________"),
        ("Student ID", "____________________________"),
        ("Submission date", "____________________________"),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(4.8)
    for row, (label, value) in zip(table.rows, details):
        cell_margins(row.cells[0], top=90, start=0, bottom=90, end=80)
        cell_margins(row.cells[1], top=90, start=0, bottom=90, end=0)
        set_font(row.cells[0].paragraphs[0].add_run(label), bold=True)
        set_font(row.cells[1].paragraphs[0].add_run(value))

    doc.add_paragraph()
    scope = doc.add_paragraph()
    set_font(scope.add_run("Report scope. "), bold=True)
    scope.add_run(
        "This report evaluates a classical Logistic Regression pipeline and a neural LSTM pipeline "
        "for classifying the SpamAssassin corpus. It explains the dataset, preprocessing, models, "
        "performance, visual evidence, cybersecurity implications and future improvements."
    )
    conclusion = doc.add_paragraph()
    set_font(conclusion.add_run("Main finding. "), bold=True)
    conclusion.add_run(
        "At the selected operating thresholds, Logistic Regression provides the stronger balance "
        "between detecting spam and limiting false alerts. The dataset contains spam and ham labels, "
        "so the result must not be presented as direct evidence of phishing-detection performance."
    )


def contents_page(doc):
    contents_heading = doc.add_paragraph("Contents", style="Heading 1")
    contents_heading.paragraph_format.page_break_before = True
    entries = [
        ("1 Cover page", "1"),
        ("2 Contents", "2"),
        ("3 Section 1 Email Security and Phishing Detection", "3"),
        ("3.1 Description of datasets used", "3"),
        ("3.2 Explanation of preprocessing techniques", "4"),
        ("3.3 Overview of implemented ML and deep learning models", "5"),
        ("3.4 Performance comparison of models", "6"),
        ("3.5 Visualizations", "7"),
        ("3.6 Discussion on findings and cybersecurity implications", "9"),
        ("3.7 Summary of key takeaways and potential improvements", "10"),
    ]
    for label, page in entries:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28 if label.startswith("3.") else 0)
        paragraph.paragraph_format.space_after = Pt(5)
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(
            Inches(6.2), alignment=WD_TAB_ALIGNMENT.RIGHT, leader=WD_TAB_LEADER.DOTS
        )
        run = paragraph.add_run(f"{label}\t{page}")
        set_font(run, size=10.5)


def section_report(doc):
    section_heading = doc.add_paragraph("3 Section 1 Email Security and Phishing Detection", style="Heading 1")
    section_heading.paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "The experiment compares two complete text-classification pipelines under the same stratified data split. "
        "The classical pipeline combines TF-IDF features with Logistic Regression, while the deep learning pipeline "
        "uses learned embeddings and an LSTM. Both models are evaluated using validation-selected thresholds that "
        "limit the validation false-positive rate to no more than one percent.",
    )

    doc.add_paragraph("3.1 Description of datasets used", style="Heading 2")
    add_paragraph(
        doc,
        "Section 1 uses the Apache SpamAssassin public corpus, specifically the easy-ham and spam archives. "
        "The source contains 3,000 email files. After parsing, normalization-aware deduplication and quality checks, "
        "2,947 records remain: 2,471 ham messages and 476 spam messages. The class distribution is therefore imbalanced, "
        "with spam forming about 16.2 percent of the retained data. [1] [2]",
    )
    add_paragraph(
        doc,
        "The shared stratified split contains 2,062 training messages, 442 validation messages and 443 test messages. "
        "The test partition contains 371 ham and 72 spam messages. Archive hashes, an extracted-file manifest hash and "
        "split membership are stored with the experiment outputs to make the fresh run auditable. [3]",
    )
    add_paragraph(
        doc,
        "A crucial limitation is the target label. The corpus distinguishes spam from ham; it does not verify whether "
        "each message is a phishing attempt. The experiment is relevant to email security, but its results should be "
        "reported as spam classification rather than direct phishing-detection performance.",
    )

    heading = add_subsection_heading(doc, "3.2 Explanation of preprocessing techniques")
    add_paragraph(
        doc,
        "Each email is read with the Python standard email parser. Attachments are excluded, the subject is combined "
        "with one preferred message body, and plain text is selected before falling back to HTML. When HTML is used, "
        "visible text is extracted with an HTML parser instead of deleting tags with a regular expression. The quality "
        "profile records 211 HTML fallbacks, three unknown-character-set fallbacks, one message without a text body and "
        "3,701 decoding replacement characters. [4]",
    )
    add_paragraph(
        doc,
        "Normalization converts text to lowercase, replaces URLs and email addresses with the stable markers urltoken "
        "and emailtoken, extracts Unicode word tokens and applies English Snowball stemming. For the TF-IDF branch, a "
        "stemmed list of 297 English stopwords is removed while no, nor and not are protected because negation can change "
        "meaning. The LSTM branch keeps stopwords so that word order and local context remain available.",
    )
    add_paragraph(
        doc,
        "Normalized representations are deduplicated before splitting. No empty normalized messages or conflicting labels "
        "were found, and the overlap check reports zero exact normalized-text matches between every pair of train, validation "
        "and test partitions. This reduces direct representation leakage, although it cannot remove every near-duplicate "
        "campaign or shared template. [4] [5]",
    )
    add_paragraph(
        doc,
        "All fitted preprocessing is learned from training data. The TF-IDF vocabulary and document frequencies are not "
        "estimated from validation or test messages, and the LSTM vocabulary is also built from the training partition only. "
        "This separation protects the final evaluation from information that would not be available during model fitting.",
    )

    heading = add_subsection_heading(doc, "3.3 Overview of implemented ML and deep learning models")
    add_paragraph(
        doc,
        "Logistic Regression is the classical model. Its input is a TF-IDF matrix containing unigrams and bigrams. The "
        "vocabulary is capped at 20,000 features, ignores terms seen in fewer than two documents or more than 98 percent "
        "of documents, and uses sublinear term frequency. Balanced class weights increase the influence of the smaller spam "
        "class, and the optimizer is allowed up to 1,000 iterations. Validation scores select threshold 0.5463 by maximizing "
        "F1 among thresholds with a validation false-positive rate no greater than one percent. [2] [6]",
    )
    add_paragraph(
        doc,
        "The deep learning model uses a training-only vocabulary with padding and unknown tokens. Every message is truncated "
        "or padded to 300 tokens. Learned 128-dimensional embeddings feed a 128-unit LSTM, followed by dropout of 0.3 and a "
        "single output logit. Weighted binary cross-entropy addresses the class imbalance, and AdamW trains the model on the "
        "CPU for six epochs. The checkpoint with the lowest validation loss is restored, and validation scores select threshold "
        "0.9772 under the same false-alert constraint. [2] [7] [8]",
    )
    add_paragraph(
        doc,
        "The models represent text differently. TF-IDF emphasizes discriminative words and adjacent word pairs but cannot "
        "model long message order. The LSTM learns embeddings and ordered context, but the available corpus may be too small "
        "to justify its greater capacity. The selected thresholds are operating points rather than calibrated deployment "
        "probabilities.",
    )

    heading = add_subsection_heading(doc, "3.4 Performance comparison of models")
    add_paragraph(
        doc,
        "The fresh test results favour Logistic Regression. It reaches F1 0.9504, compared with 0.7797 for the LSTM, and "
        "retains much more spam recall while producing only two false alerts. The LSTM achieves perfect observed precision "
        "because it produces no false positives, but it detects only 46 of 72 spam messages and misses 26. [6] [7]",
    )
    add_metrics_table(
        doc,
        [
            ["Logistic Regression", "0.9842", "0.9710", "0.9306", "0.9504", "0.9996", "0.9977"],
            ["LSTM", "0.9413", "1.0000", "0.6389", "0.7797", "0.9728", "0.9444"],
        ],
    )
    add_paragraph(doc, "Test confusion-matrix counts", bold_lead="Test confusion-matrix counts")
    add_confusion_table(doc)
    add_paragraph(
        doc,
        "Ranking metrics lead to the same overall conclusion. Logistic Regression obtains ROC AUC 0.9996 and average "
        "precision 0.9977, compared with 0.9728 and 0.9444 for the LSTM. The LSTM training loss continues to fall after "
        "epoch three while validation loss rises before partially recovering. Restoring the epoch-three checkpoint is "
        "therefore justified and suggests that later training increasingly fits the training sample rather than improving "
        "generalization. [6] [7] [8]",
    )

    heading = add_subsection_heading(doc, "3.5 Visualizations")
    visualization_intro = add_paragraph(
        doc,
        "The figures below show the class and message-length distributions, test confusion matrices, ROC and precision-recall "
        "curves, and the LSTM learning history. Together they reveal both ranking quality and the operational effect of the "
        "selected thresholds.",
    )
    heading.paragraph_format.left_indent = Inches(0.15)
    visualization_intro.paragraph_format.left_indent = Inches(0.15)
    add_figure(
        doc,
        ROOT / "section_01_results/reports/figures/data-exploration.png",
        "Class balance and normalized message-length distributions for the retained email corpus.",
        1,
    )
    add_figure(
        doc,
        ROOT / "section_01_results/reports/figures/classic-evaluation.png",
        "Logistic Regression confusion matrix, ROC curve and precision-recall curve on the test set.",
        2,
    )
    add_figure(
        doc,
        ROOT / "section_01_results/reports/figures/lstm-evaluation.png",
        "LSTM confusion matrix, ROC curve and precision-recall curve on the test set.",
        3,
        page_break_before=True,
    )
    add_figure(
        doc,
        ROOT / "section_01_results/reports/figures/lstm-training-history.png",
        "LSTM training and validation loss across six epochs.",
        4,
    )

    heading = add_subsection_heading(doc, "3.6 Discussion on findings and cybersecurity implications")
    add_paragraph(
        doc,
        "The comparison shows that the more complex neural architecture is not automatically the better security model. "
        "For this corpus and operating policy, the sparse linear representation separates the classes more effectively and "
        "provides a stronger balance between legitimate-email protection and spam detection. This is operationally important "
        "because false positives can hide or delay legitimate communication, while false negatives allow malicious or unwanted "
        "messages to reach users.",
    )
    add_paragraph(
        doc,
        "The LSTM's perfect observed precision must be interpreted with its low recall. It avoids false alerts by using a high "
        "threshold, but misses more than one third of the spam messages in the test set. A security team should therefore inspect "
        "precision, recall, confusion counts and threshold policy together instead of relying on accuracy or one headline metric.",
    )
    add_paragraph(
        doc,
        "The experiment also has limits that prevent deployment claims. The corpus is historical, the positive class is spam "
        "rather than verified phishing, and related campaigns may share templates even after exact normalized duplicates are "
        "removed. URL and email replacement can improve generalization but also removes identities and structures that may be "
        "useful indicators. Neither model has been evaluated against recent phishing techniques, adversarial wording, multilingual "
        "messages or a live organization's acceptable false-alert workload.",
    )

    heading = add_subsection_heading(doc, "3.7 Summary of key takeaways and potential improvements")
    add_bullets(
        doc,
        [
            "Logistic Regression is the stronger model in the fresh run, with F1 0.9504 and only two false positives.",
            "The LSTM ranks messages well but its selected threshold sacrifices recall, missing 26 of 72 test spam messages.",
            "Training-only fitting and zero exact normalized-text overlap strengthen the evaluation design.",
            "The corpus supports spam classification claims, not direct claims about phishing detection or prevention.",
            "A focused error analysis should inspect the Logistic Regression false positives and false negatives and the LSTM false negatives using redacted examples.",
            "Future evaluation should use recent and independently labelled phishing data, grouped or time-based splits, repeated seeds and uncertainty intervals.",
            "Thresholds should be selected against an explicit operational false-positive budget, followed by an untouched final evaluation set.",
            "Further improvements could compare calibrated probabilities, stronger linear baselines, pretrained language models and ablations of stemming, stopwords, token replacement and sequence length.",
        ],
    )
    add_paragraph(
        doc,
        "Overall, the experiment demonstrates a clear and reproducible email text-classification workflow. Its strongest result "
        "is the classical TF-IDF and Logistic Regression pipeline, but the evidence remains an educational benchmark rather than "
        "proof of a production phishing-defence system.",
    )

    evidence_heading = doc.add_paragraph("Evidence sources", style="Heading 2")
    keep_with_next(evidence_heading)
    evidence = [
        ("[1] SpamAssassin corpus record", "docs/research/spamassassin-public-corpus.md"),
        ("[2] Section 1 notebook", "notebooks/01_phishing/section-01-email-security.ipynb"),
        ("[3] Fresh run summary", "section_01_results/reports/run-summary.json"),
        ("[4] Text quality profile", "section_01_results/processed/text-profile.json"),
        ("[5] Representation overlap checks", "section_01_results/processed/representation-overlap.json"),
        ("[6] Logistic Regression metrics", "section_01_results/reports/metrics/classic.json"),
        ("[7] LSTM metrics", "section_01_results/reports/metrics/lstm.json"),
        ("[8] LSTM training history", "section_01_results/reports/metrics/lstm-training-history.csv"),
    ]
    for label, path in evidence:
        paragraph = doc.add_paragraph(style="List Bullet")
        add_hyperlink(paragraph, label, path)


def build():
    doc = Document()
    configure_document(doc)
    cover_page(doc)
    contents_page(doc)
    section_report(doc)

    properties = doc.core_properties
    properties.title = "Section 1 Email Security and Phishing Detection"
    properties.subject = "Machine Learning in Cyber assignment report"
    properties.keywords = "email security, spam detection, phishing, logistic regression, LSTM"
    properties.comments = "Prepared from the updated Section 1 notebook and fresh exported evidence."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
